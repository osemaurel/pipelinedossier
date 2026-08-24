"""Introspection du modèle Excel.

Le classeur téléversé est la source de vérité : rien de sa structure n'est codé
en dur ici. Ce module en déduit, par lecture seule, les feuilles, la ligne
d'en-tête réelle, la ligne d'exemple, les colonnes, leurs formules modèles,
leurs listes de valeurs autorisées et leurs contraintes de longueur.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.core.logging import get_logger

logger = get_logger(__name__)

# Références de cellules relatives / absolues dans une formule.
_CELL_REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
_RANGE_REF = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\-]+))!\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
_LEN_GT = re.compile(r"LEN\(\s*\$?([A-Z]{1,3})\$?\d+\s*\)\s*>\s*(\d+)", re.IGNORECASE)
_LEN_LT = re.compile(r"LEN\(\s*\$?([A-Z]{1,3})\$?\d+\s*\)\s*<\s*(\d+)", re.IGNORECASE)
_HEADER_MAX = re.compile(r"(\d+)\s*caract", re.IGNORECASE)
_HEADER_RANGE = re.compile(r"(\d+)\s*(?:à|a|-)\s*(\d+)\s*caract", re.IGNORECASE)
_HEADER_ITEMS_MAX = re.compile(r"(\d+)\s*max", re.IGNORECASE)

# Marqueur précis plutôt que le simple mot « exemple » : une ligne de données
# réelle peut contenir ce mot, et la confondre avec l'exemple ferait ignorer un
# profil déjà saisi lors de la reprise de numérotation.
EXAMPLE_MARKERS = ("ligne d exemple",)


def normalize(text: object) -> str:
    """Clé de comparaison : sans accents, sans ponctuation, minuscule."""
    if text is None:
        return ""
    raw = str(text)
    decomposed = unicodedata.normalize("NFKD", raw)
    stripped = "".join(c for c in decomposed if not unicodedata.combining(c))
    stripped = stripped.replace("’", "'").replace("—", " ").replace("–", " ")
    stripped = re.sub(r"\(.*?\)", " ", stripped)
    stripped = re.sub(r"[^A-Za-z0-9]+", " ", stripped)
    return " ".join(stripped.lower().split())


@dataclass(slots=True)
class ColumnSchema:
    index: int
    letter: str
    header: str
    key: str
    formula_template: str | None = None
    allowed_values: list[str] = field(default_factory=list)
    number_format: str = "General"
    min_length: int | None = None
    max_length: int | None = None
    max_items: int | None = None
    is_date: bool = False

    @property
    def is_formula(self) -> bool:
        return self.formula_template is not None


@dataclass(slots=True)
class SheetSchema:
    name: str
    header_row: int
    data_start_row: int
    data_end_row: int
    example_row: int | None
    columns: list[ColumnSchema]
    # Lignes déjà renseignées par un humain, ligne d'exemple exclue.
    filled_rows: list[int] = field(default_factory=list)
    existing_codes: list[str] = field(default_factory=list)

    @property
    def capacity(self) -> int:
        return self.data_end_row - self.data_start_row + 1

    @property
    def first_free_row(self) -> int:
        """Première ligne libre : on écrit à la suite, jamais par-dessus."""
        return max(self.filled_rows) + 1 if self.filled_rows else self.data_start_row

    def next_number(self, pattern: re.Pattern[str]) -> int:
        """Numéro suivant le plus grand déjà attribué dans la colonne des codes."""
        numbers = [
            int(match.group(1))
            for code in self.existing_codes
            if (match := pattern.match(code))
        ]
        return max(numbers) + 1 if numbers else 1

    def by_key(self, key: str) -> ColumnSchema | None:
        for column in self.columns:
            if column.key == key:
                return column
        return None

    def find(self, *fragments: str) -> ColumnSchema | None:
        """Première colonne dont la clé contient tous les fragments donnés."""
        for column in self.columns:
            if all(fragment in column.key for fragment in fragments):
                return column
        return None


@dataclass(slots=True)
class WorkbookSchema:
    path: Path
    sheet_names: list[str]
    femmes: SheetSchema
    agents: SheetSchema
    photos: SheetSchema | None
    listes: dict[str, list[str]]
    readme: list[tuple[str, str]]

    def text_constraints(self) -> dict[str, tuple[int | None, int | None]]:
        return {
            column.key: (column.min_length, column.max_length)
            for column in self.femmes.columns
            if column.max_length is not None
        }


def _is_inside_multicell_merge(ws: Worksheet, row: int, col: int) -> bool:
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            if merged.max_col > merged.min_col or merged.max_row > merged.min_row:
                return True
    return False


def _detect_header_row(ws: Worksheet, limit: int = 6) -> int:
    """La ligne d'en-tête est la plus dense en libellés hors cellules fusionnées."""
    best_row, best_score = 1, -1
    for row in range(1, min(limit, ws.max_row) + 1):
        score = 0
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row=row, column=col).value
            if not isinstance(value, str) or not value.strip():
                continue
            if _is_inside_multicell_merge(ws, row, col):
                continue
            score += 1
        if score > best_score:
            best_row, best_score = row, score
    return best_row


def _detect_example_row(ws: Worksheet, header_row: int) -> int | None:
    """Repère la ligne d'exemple à son marqueur « LIGNE D'EXEMPLE »."""
    for row in range(header_row + 1, min(header_row + 4, ws.max_row) + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row=row, column=col).value
            if isinstance(value, str) and any(m in normalize(value) for m in EXAMPLE_MARKERS):
                return row
    # Pas de repli sur la couleur de fond : une ligne d'exemple effacée puis
    # réutilisée pour de vraies données garde sa teinte, et la prendre pour un
    # exemple rendrait ce profil invisible aux contrôles de doublon. Se tromper
    # dans l'autre sens est sans danger : on écrit simplement à la suite.
    return None


def _detect_data_end(ws: Worksheet, header_row: int) -> int:
    """Dernière ligne préparée du modèle.

    `max_row` reflète l'étendue déclarée du modèle (styles et validations
    pré-posés), pas seulement les cellules remplies : c'est bien la capacité
    prévue par l'auteur du classeur.
    """
    return max(ws.max_row, header_row + 1)


def _formula_template(formula: str, source_row: int) -> str:
    """Remplace les références de ligne *relatives* égales à `source_row` par `{row}`."""

    def replace(match: re.Match[str]) -> str:
        col_abs, col, row_abs, row = match.groups()
        if row_abs == "$" or int(row) != source_row:
            return match.group(0)
        return f"{col_abs}{col}{row_abs}{{row}}"

    return _CELL_REF.sub(replace, formula)


def retarget_ranges(formula: str, sheet_name: str, start_row: int, end_row: int) -> str:
    """Réaligne les plages `Feuille!$X$a:$X$b` sur l'étendue réelle des données.

    Le modèle fige ses plages sur la capacité initiale (ex. `Femmes!$B$4:$B$153`).
    Réécrire les bornes garde la formule vivante quand les données commencent ou
    finissent ailleurs, plutôt que de la remplacer par une valeur statique.
    """
    target = normalize(sheet_name)

    def replace(match: re.Match[str]) -> str:
        quoted, plain, col1, _row1, col2, _row2 = match.groups()
        name = quoted or plain
        if normalize(name) != target:
            return match.group(0)
        prefix = f"'{name}'!" if quoted else f"{name}!"
        return f"{prefix}${col1}${start_row}:${col2}${end_row}"

    return _RANGE_REF.sub(replace, formula)


def _resolve_list_values(wb: Workbook, formula1: str | None) -> list[str]:
    """Résout la source d'une liste déroulante en valeurs concrètes."""
    if not formula1:
        return []
    raw = formula1.strip().lstrip("=")
    if raw.startswith('"') and raw.endswith('"'):
        return [part.strip() for part in raw[1:-1].split(",") if part.strip()]

    match = re.match(r"(?:'([^']+)'|([A-Za-z0-9_\-]+))!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)", raw)
    if not match:
        return []
    sheet_name = match.group(1) or match.group(2)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(match.group(3).replace("$", ""))
    values: list[str] = []
    for row in range(min_row or 1, (max_row or min_row or 1) + 1):
        for col in range(min_col or 1, (max_col or min_col or 1) + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None and str(value).strip():
                values.append(str(value).strip())
    return values


def _validation_map(wb: Workbook, ws: Worksheet) -> dict[int, list[str]]:
    """Colonne (index) -> valeurs autorisées, d'après les validations de données."""
    mapping: dict[int, list[str]] = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        values = _resolve_list_values(wb, str(dv.formula1) if dv.formula1 else None)
        if not values:
            continue
        for cell_range in dv.sqref.ranges:
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                mapping.setdefault(col, values)
    return mapping


def _length_constraints(ws: Worksheet, data_row: int) -> dict[int, tuple[int | None, int | None]]:
    """Extrait min/max de longueur depuis les formules de contrôle `LEN(...)`.

    Les formules du modèle sont plus précises que les en-têtes : l'en-tête
    « Accroche (120 caractères max) » tait le minimum de 40 que la formule impose.
    """
    constraints: dict[int, tuple[int | None, int | None]] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=data_row, column=col).value
        if not isinstance(value, str) or not value.startswith("="):
            continue
        maxima = _LEN_GT.findall(value)
        minima = _LEN_LT.findall(value)
        for letter, bound in maxima:
            index = _letter_to_index(letter)
            low = next((int(b) for lt, b in minima if lt == letter), None)
            constraints[index] = (low, int(bound))
    return constraints


def _letter_to_index(letter: str) -> int:
    index = 0
    for char in letter:
        index = index * 26 + (ord(char.upper()) - ord("A") + 1)
    return index


def _header_constraints(header: str) -> tuple[int | None, int | None, int | None]:
    """Contraintes lisibles dans l'en-tête (repli si aucune formule de contrôle)."""
    ranged = _HEADER_RANGE.search(header)
    if ranged:
        low, high = int(ranged.group(1)), int(ranged.group(2))
        return min(low, high), max(low, high), None
    single = _HEADER_MAX.search(header)
    items = _HEADER_ITEMS_MAX.search(header)
    max_items = int(items.group(1)) if items and not single else None
    if single:
        return None, int(single.group(1)), max_items
    return None, None, max_items


def _read_sheet(wb: Workbook, ws: Worksheet) -> SheetSchema:
    header_row = _detect_header_row(ws)
    example_row = _detect_example_row(ws, header_row)
    data_start = example_row if example_row is not None else header_row + 1
    columns_count = ws.max_column
    data_end = _detect_data_end(ws, header_row)
    validations = _validation_map(wb, ws)
    lengths = _length_constraints(ws, data_start)

    columns: list[ColumnSchema] = []
    for index in range(1, columns_count + 1):
        header = ws.cell(row=header_row, column=index).value
        if header is None or not str(header).strip():
            continue
        header_text = str(header).strip()
        sample = ws.cell(row=data_start, column=index)
        formula = None
        if isinstance(sample.value, str) and sample.value.startswith("="):
            formula = _formula_template(sample.value, data_start)

        min_len, max_len = lengths.get(index, (None, None))
        h_min, h_max, h_items = _header_constraints(header_text)
        columns.append(
            ColumnSchema(
                index=index,
                letter=get_column_letter(index),
                header=header_text,
                key=normalize(header_text),
                formula_template=formula,
                allowed_values=validations.get(index, []),
                number_format=sample.number_format or "General",
                min_length=min_len if min_len is not None else h_min,
                max_length=max_len if max_len is not None else h_max,
                max_items=h_items,
                is_date="YY" in (sample.number_format or "").upper(),
            )
        )

    filled_rows, existing_codes = _read_existing(ws, data_start, data_end, example_row)
    return SheetSchema(
        name=ws.title,
        header_row=header_row,
        data_start_row=data_start,
        data_end_row=data_end,
        example_row=example_row,
        columns=columns,
        filled_rows=filled_rows,
        existing_codes=existing_codes,
    )


def _read_existing(
    ws: Worksheet, start: int, end: int, example_row: int | None
) -> tuple[list[int], list[str]]:
    """Lignes déjà saisies et codes qu'elles portent.

    La ligne d'exemple est écartée : elle porte un code d'illustration
    (« PAL-0001 ») qui fausserait la reprise de la numérotation.
    """
    rows: list[int] = []
    codes: list[str] = []
    for row in range(start, min(end, ws.max_row) + 1):
        if row == example_row:
            continue
        value = ws.cell(row=row, column=1).value
        if value is None or not str(value).strip():
            continue
        rows.append(row)
        codes.append(str(value).strip())
    return rows, codes


def _read_listes(ws: Worksheet) -> dict[str, list[str]]:
    listes: dict[str, list[str]] = {}
    for col in range(1, ws.max_column + 1):
        title = ws.cell(row=1, column=col).value
        if title is None:
            continue
        values = [
            str(ws.cell(row=row, column=col).value).strip()
            for row in range(2, ws.max_row + 1)
            if ws.cell(row=row, column=col).value is not None
        ]
        if values:
            listes[str(title).strip()] = values
    return listes


def _read_readme(ws: Worksheet) -> list[tuple[str, str]]:
    entries: list[tuple[str, str]] = []
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row=row, column=2).value
        text = ws.cell(row=row, column=3).value
        if label is None and text is None:
            continue
        entries.append((str(label or "").strip(), str(text or "").strip()))
    return entries


def _match_sheet(wb: Workbook, *fragments: str) -> Worksheet | None:
    for name in wb.sheetnames:
        key = normalize(name)
        if any(fragment in key for fragment in fragments):
            return wb[name]
    return None


def introspect(path: Path) -> WorkbookSchema:
    """Analyse le modèle et renvoie son schéma. N'écrit jamais dans le fichier."""
    wb = load_workbook(path, data_only=False)
    try:
        femmes_ws = _match_sheet(wb, "femme", "profil")
        agents_ws = _match_sheet(wb, "agent")
        photos_ws = _match_sheet(wb, "photo")
        listes_ws = _match_sheet(wb, "liste")
        readme_ws = _match_sheet(wb, "lisez", "readme", "lisezmoi")

        if femmes_ws is None or agents_ws is None:
            missing = "Femmes" if femmes_ws is None else "Agents"
            raise ValueError(
                f"Modèle invalide : feuille « {missing} » introuvable. "
                f"Feuilles présentes : {', '.join(wb.sheetnames)}."
            )

        schema = WorkbookSchema(
            path=path,
            sheet_names=list(wb.sheetnames),
            femmes=_read_sheet(wb, femmes_ws),
            agents=_read_sheet(wb, agents_ws),
            photos=_read_sheet(wb, photos_ws) if photos_ws is not None else None,
            listes=_read_listes(listes_ws) if listes_ws is not None else {},
            readme=_read_readme(readme_ws) if readme_ws is not None else [],
        )
    finally:
        wb.close()

    logger.info(
        "Modèle analysé : %d feuilles, Femmes=%d colonnes (en-tête ligne %d, données %d→%d), "
        "Agents=%d colonnes, Photos=%d colonnes, %d listes",
        len(schema.sheet_names),
        len(schema.femmes.columns),
        schema.femmes.header_row,
        schema.femmes.data_start_row,
        schema.femmes.data_end_row,
        len(schema.agents.columns),
        len(schema.photos.columns) if schema.photos else 0,
        len(schema.listes),
    )
    return schema
