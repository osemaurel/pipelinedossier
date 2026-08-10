"""Analyse détaillée d'un modèle de collecte Excel (diagnostic, lecture seule)."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def describe(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    print(f"=== {path.name} ===")
    print("Feuilles :", wb.sheetnames)
    print("Noms définis :", list(wb.defined_names))
    for ws in wb.worksheets:
        print(f"\n--- Feuille « {ws.title} » ---")
        print(f"dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} "
              f"freeze={ws.freeze_panes} state={ws.sheet_state}")
        print(f"tables={list(ws.tables)} merged={[str(r) for r in ws.merged_cells.ranges][:10]}")
        print(f"cond_formatting={len(list(ws.conditional_formatting))}")

        widths = {k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}
        print("largeurs:", widths)

        print("Lignes (12 premières, 14 premières colonnes) :")
        for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row),
                                max_col=min(14, ws.max_column)):
            cells = []
            for c in row:
                if c.value is None:
                    continue
                cells.append(f"{c.coordinate}={c.value!r}")
            if cells:
                print("   ", " | ".join(cells)[:400])

        formulas = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append((c.coordinate, c.value))
        print(f"formules: {len(formulas)}")
        seen: set[str] = set()
        for coord, f in formulas:
            col = "".join(ch for ch in coord if ch.isalpha())
            if col in seen:
                continue
            seen.add(col)
            print(f"    {coord}: {f[:260]}")

        dvs = ws.data_validations.dataValidation
        print(f"validations: {len(dvs)}")
        for dv in dvs:
            print(f"    type={dv.type} op={dv.operator} f1={str(dv.formula1)[:160]} "
                  f"f2={dv.formula2} ranges={str(dv.sqref)[:120]}")


if __name__ == "__main__":
    describe(Path(sys.argv[1] if len(sys.argv) > 1 else "palabdossiercollecte.xlsx"))
