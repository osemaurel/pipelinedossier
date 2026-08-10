"""Vérifie un dossier produit : intégrité du modèle, données, formules, photos, ZIP.

Usage : python -m tests.verify_output <dossier_de_sortie>
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from backend.services.excel_introspect import introspect

TEMPLATE = Path("palabdossiercollecte.xlsx")


def verify(output_dir: Path) -> int:
    excel = next(output_dir.glob("*.xlsx"))
    archive = output_dir / "palab_dossier_final.zip"
    failures: list[str] = []

    def check(label: str, condition: bool, detail: str = "") -> None:
        print(f"{'✓' if condition else '✗'} {label}{(' — ' + detail) if detail else ''}")
        if not condition:
            failures.append(label)

    reference = introspect(TEMPLATE)
    produced = load_workbook(excel, data_only=False)
    original = load_workbook(TEMPLATE, data_only=False)

    print("\n--- Structure ---")
    check("feuilles identiques au modèle", produced.sheetnames == original.sheetnames,
          str(produced.sheetnames))

    femmes = produced[reference.femmes.name]
    agents = produced[reference.agents.name]
    photos = produced[reference.photos.name]
    orig_femmes = original[reference.femmes.name]

    check("en-têtes de la feuille Femmes inchangés",
          [femmes.cell(row=2, column=c).value for c in range(1, 46)]
          == [orig_femmes.cell(row=2, column=c).value for c in range(1, 46)])
    check("bandes fusionnées conservées",
          len(femmes.merged_cells.ranges) == len(orig_femmes.merged_cells.ranges),
          f"{len(femmes.merged_cells.ranges)} plages")
    check("largeurs de colonnes conservées",
          femmes.column_dimensions["AK"].width == orig_femmes.column_dimensions["AK"].width)
    check("listes déroulantes conservées",
          len(femmes.data_validations.dataValidation)
          == len(orig_femmes.data_validations.dataValidation),
          f"{len(femmes.data_validations.dataValidation)} validations")
    check("volets figés conservés", femmes.freeze_panes == orig_femmes.freeze_panes)

    print("\n--- Données ---")
    start = reference.femmes.data_start_row
    codes = []
    row = start
    while femmes.cell(row=row, column=1).value:
        codes.append(str(femmes.cell(row=row, column=1).value))
        row += 1
    written = len(codes)
    check("profils écrits", written > 0, f"{written} lignes à partir de la ligne {start}")
    check("codes femme uniques et bien formés",
          len(set(codes)) == written and all(c.startswith("PAL-") for c in codes))
    check("ligne d'exemple neutralisée",
          "EXEMPLE" not in str(femmes.cell(row=start, column=45).value or "").upper(),
          str(femmes.cell(row=start, column=45).value)[:60])

    print("\n--- Formules ---")
    for letter, col in (("E", 5), ("AO", 41), ("AP", 42), ("AQ", 43)):
        values = [femmes.cell(row=r, column=col).value for r in range(start, start + written)]
        ok = all(isinstance(v, str) and v.startswith("=") for v in values)
        check(f"formule Femmes!{letter} présente sur toutes les lignes de données", ok,
              str(values[0])[:60])

    agent_formula = str(agents.cell(row=reference.agents.data_start_row, column=11).value)
    check("formule Agents!K présente", agent_formula.startswith("="), agent_formula[:70])
    check("plage de la formule Agents!K réalignée sur les données",
          f"$B${start}:" in agent_formula, agent_formula[:70])

    print("\n--- Colonnes de vérification ---")
    for letter, col in (("K", 11), ("L", 12), ("M", 13), ("N", 14), ("O", 15)):
        value = str(femmes.cell(row=start, column=col).value or "")
        check(f"colonne {letter} non fabriquée", "démonstration" in value, value)
    check("statut du dossier au plus bas",
          str(femmes.cell(row=start, column=44).value) == "À compléter",
          str(femmes.cell(row=start, column=44).value))
    check("notes internes marquées fictives",
          "FICTIF" in str(femmes.cell(row=start, column=45).value or "").upper())

    print("\n--- Contacts fictifs ---")
    emails = [str(femmes.cell(row=r, column=9).value) for r in range(start, start + written)]
    check("e-mails en domaine réservé example.test",
          all(e.endswith("@example.test") for e in emails), emails[0])

    print("\n--- Longueurs de texte (bornes du modèle) ---")
    for letter, col, low, high in (("AJ", 36, 40, 120), ("AK", 37, 600, 1200), ("AL", 38, 200, 400)):
        lengths = [len(str(femmes.cell(row=r, column=col).value or ""))
                   for r in range(start, start + written)]
        check(f"colonne {letter} dans [{low}, {high}]",
              all(low <= n <= high for n in lengths), f"longueurs={lengths}")

    print("\n--- Photos ---")
    photo_start = reference.photos.data_start_row
    rows = []
    r = photo_start
    while photos.cell(row=r, column=1).value:
        rows.append((str(photos.cell(row=r, column=1).value),
                     str(photos.cell(row=r, column=2).value),
                     photos.cell(row=r, column=3).value))
        r += 1
    check("lignes Photos écrites", len(rows) > 0, f"{len(rows)} lignes")
    check("codes Photos tous présents dans Femmes",
          all(code in set(codes) for code, _, _ in rows))

    on_disk = {p.name for p in (output_dir.parent.parent / "temp").rglob("*.png")}
    in_zip: set[str] = set()
    with zipfile.ZipFile(archive) as zf:
        names = zf.namelist()
        in_zip = {Path(n).name for n in names if n.endswith(".png")}
        check("ZIP contient le classeur", any(n.endswith(".xlsx") for n in names))
        check("ZIP contient README.txt", any(n.endswith("README.txt") for n in names))
        check("ZIP contient le rapport",
              any(n.endswith("rapport_validation.txt") for n in names))
        check("ZIP valide", zf.testzip() is None)

    referenced = {name for _, name, _ in rows}
    check("chaque fichier référencé dans Excel est présent dans le ZIP",
          referenced <= in_zip, f"manquants={sorted(referenced - in_zip)}")
    check("le nombre de photos annoncé correspond aux fichiers",
          all(
              femmes.cell(row=start + i, column=39).value
              == sum(1 for code, _, _ in rows if code == codes[i])
              for i in range(written)
          ))

    print("\n--- Provenance des images ---")
    with zipfile.ZipFile(archive) as zf:
        sample = next(n for n in zf.namelist() if n.endswith(".png"))
        blob = zf.read(sample)
    check("métadonnées PNG de provenance IA", b"Palab Dossier Generator" in blob, sample)

    print("\n--- Modèle source ---")
    check("le fichier modèle est resté intact",
          TEMPLATE.stat().st_size == 47577 and orig_femmes.cell(row=3, column=1).value == "PAL-0001")

    print(f"\n{'TOUT EST CONFORME' if not failures else f'{len(failures)} ÉCHEC(S) : ' + ', '.join(failures)}")
    return 0 if not failures else 1


if __name__ == "__main__":
    raise SystemExit(verify(Path(sys.argv[1])))
