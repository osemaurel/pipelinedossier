"""Tests unitaires des briques qui ne dépendent pas d'un appel réseau."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from backend.services.excel_introspect import introspect, normalize, retarget_ranges
from backend.services.field_policy import (
    FillMode,
    build_agent_specs,
    build_femme_specs,
    spec_by_name,
)
from backend.services.image_generator import stamp_provenance
from backend.services.validation_service import CODE_FEMME
from backend.services.profile_generator import (
    birth_window,
    build_json_schema,
    demo_email,
    demo_phone,
    sentence_safe_trim,
)

TEMPLATE = Path("palabdossiercollecte.xlsx")


@pytest.fixture(scope="module")
def schema():
    return introspect(TEMPLATE)


def test_normalize_strips_accents_and_parentheses():
    assert normalize("Accroche (120 caractères max)") == "accroche"
    assert normalize("Niveau d'études") == "niveau d etudes"
    assert normalize("Âge recherché — min") == "age recherche min"


def test_introspection_finds_real_header_row(schema):
    """La feuille Femmes a une ligne de bandes fusionnées au-dessus des en-têtes."""
    assert schema.femmes.header_row == 2
    assert schema.femmes.example_row == 3
    assert len(schema.femmes.columns) == 45
    assert schema.femmes.by_key("code femme") is not None


def test_length_bounds_come_from_control_formulas(schema):
    """Les formules sont plus précises que les en-têtes : elles portent le minimum."""
    constraints = schema.text_constraints()
    assert constraints["accroche"] == (40, 120)
    assert constraints["presentation"] == (600, 1200)
    assert constraints["ce que je recherche"] == (200, 400)


def test_dropdown_values_are_resolved_from_listes_sheet(schema):
    situation = schema.femmes.find("situation")
    assert situation is not None
    assert situation.allowed_values == ["Célibataire", "Divorcée", "Veuve", "Séparée"]


def test_formula_templates_are_row_parameterised(schema):
    age = schema.femmes.by_key("age")
    assert age is not None
    assert age.formula_template == '=IF(D{row}="","",DATEDIF(D{row},TODAY(),"Y"))'


def test_retarget_ranges_only_touches_the_named_sheet():
    formula = '=IF(A2="","",COUNTIF(Femmes!$B$4:$B$153,A2))'
    assert retarget_ranges(formula, "Femmes", 3, 162) == (
        '=IF(A2="","",COUNTIF(Femmes!$B$3:$B$162,A2))'
    )
    assert retarget_ranges(formula, "Photos", 3, 162) == formula


def test_verification_columns_are_never_generated(schema):
    specs = build_femme_specs(schema.femmes)
    withheld = {spec.name for spec in specs if spec.mode is FillMode.WITHHELD}
    assert withheld == {
        "piece_identite", "numero_piece", "contrat_mandat",
        "date_signature", "consentement_photos",
    }
    assert spec_by_name(specs, "statut_dossier").mode is FillMode.STATUS
    assert spec_by_name(specs, "email").mode is FillMode.CONTACT


def test_agent_columns_are_classified(schema):
    specs = build_agent_specs(schema.agents)
    assert spec_by_name(specs, "code_agent").mode is FillMode.ASSIGNED
    assert spec_by_name(specs, "nom_agence").mode is FillMode.GENERATED
    assert any(spec.mode is FillMode.COMPUTED for spec in specs)


def test_json_schema_is_strict_and_carries_enums(schema):
    specs = build_femme_specs(schema.femmes)
    built = build_json_schema(specs, "profils")
    item = built["properties"]["profils"]["items"]
    assert item["additionalProperties"] is False
    assert set(item["required"]) == set(item["properties"])
    assert "piece_identite" not in item["properties"]
    assert item["properties"]["situation"]["enum"] == ["Célibataire", "Divorcée", "Veuve", "Séparée"]
    assert item["properties"]["taille_cm"]["type"] == "integer"


def test_json_schema_filters_narrow_but_never_widen(schema):
    specs = build_femme_specs(schema.femmes)
    built = build_json_schema(
        specs, "profils", {"situation": ["Divorcée", "Valeur inexistante"]}
    )
    assert built["properties"]["profils"]["items"]["properties"]["situation"]["enum"] == ["Divorcée"]


def test_birth_window_matches_requested_ages():
    reference = date(2026, 8, 10)
    oldest, youngest = birth_window(25, 45, reference)
    for born in (oldest, youngest):
        age = reference.year - born.year - ((reference.month, reference.day) < (born.month, born.day))
        assert 25 <= age <= 45


def test_demo_contacts_use_reserved_ranges():
    assert demo_email("Aminata", "PAL-0001").endswith("@example.test")
    assert demo_phone(7).startswith("+99")


def test_sentence_safe_trim_never_cuts_mid_sentence():
    text = "Première phrase complète. Deuxième phrase un peu plus longue. Troisième phrase."
    trimmed = sentence_safe_trim(text, 40)
    assert len(trimmed) <= 40
    assert trimmed.endswith(".")
    assert trimmed == "Première phrase complète."


def test_sentence_safe_trim_keeps_short_text_untouched():
    assert sentence_safe_trim("Court.", 100) == "Court."


def test_png_provenance_chunk_is_inserted():
    png = (
        b"\x89PNG\r\n\x1a\n"
        + b"\x00\x00\x00\x0dIHDR" + b"\x00" * 13 + b"\x00\x00\x00\x00"
        + b"\x00\x00\x00\x00IEND\xae B`\x82"
    )
    stamped = stamp_provenance(png)
    assert b"Palab Dossier Generator" in stamped
    assert stamped.startswith(b"\x89PNG\r\n\x1a\n")
    assert b"IHDR" in stamped[:30]


def test_stamp_provenance_ignores_non_png():
    assert stamp_provenance(b"pas une image") == b"pas une image"


def _avatar_ctx(code: str, variant: int):
    from backend.prompts.avatar_prompt import AvatarContext, AvatarStyle

    return AvatarContext(
        code_femme=code, age=31, ville="Abidjan", pays="Côte d'Ivoire",
        profession="Infirmière", yeux="Marron", cheveux="Noir", variant_index=variant,
        style=AvatarStyle.PHOTO, taille_cm=168, poids_kg=62, nationalite="Ivoirienne",
    )


def _sheet_of(prompt: str) -> str:
    """La fiche personnage est le deuxième bloc du prompt."""
    return prompt.split("\n\n")[1]


def test_character_sheet_is_identical_across_a_profiles_photos():
    """Le visage dérivait d'une image à l'autre : la fiche doit être figée."""
    from backend.prompts.avatar_prompt import build_avatar_prompt

    sheets = {_sheet_of(build_avatar_prompt(_avatar_ctx("PAL-0001", i))) for i in range(8)}
    assert len(sheets) == 1, "la description physique change d'une photo à l'autre"


def test_character_sheet_pins_complexion_and_build():
    """Sans carnation ni corpulence, la même femme sortait claire puis foncée."""
    from backend.prompts.avatar_prompt import build_character_sheet

    sheet = build_character_sheet(_avatar_ctx("PAL-0001", 0))
    assert "Carnation :" in sheet and "peau" in sheet
    assert "Corpulence :" in sheet and "168 cm" in sheet
    assert "Visage :" in sheet
    assert "Ne modifie ni la couleur de peau" in sheet


def test_build_follows_height_and_weight():
    """La corpulence vient de la taille et du poids du dossier, pas du hasard."""
    from backend.prompts.avatar_prompt import AvatarContext, build_character_sheet

    def sheet(taille: int, poids: int) -> str:
        return build_character_sheet(AvatarContext(
            code_femme="PAL-0001", age=31, ville="Abidjan", pays="Côte d'Ivoire",
            profession="Infirmière", yeux="Marron", cheveux="Noir", variant_index=0,
            taille_cm=taille, poids_kg=poids, nationalite="Ivoirienne",
        ))

    assert "mince" in sheet(175, 52)
    assert "moyenne" in sheet(168, 62)
    assert "forte" in sheet(160, 85)


def test_eight_photos_are_not_eight_selfies():
    """Demander huit photos donnait huit selfies : les prises doivent alterner."""
    from backend.prompts.avatar_prompt import build_scene

    scenes = [build_scene(_avatar_ctx("PAL-0001", i)) for i in range(8)]
    selfies = [s for s in scenes if s.lower().startswith("selfie")]
    assert len(set(scenes)) == 8, "deux photos partagent la même scène"
    assert 2 <= len(selfies) <= 5, f"{len(selfies)} selfies sur 8"


def test_lighting_matches_the_setting():
    """Une scène de rue éclairée au plafonnier trahit immédiatement l'image."""
    from backend.prompts.avatar_prompt import (
        CANDID_SCENES, LIGHTS_INDOOR, LIGHTS_OUTDOOR, build_scene,
    )

    indoor = {light.lower() for light in LIGHTS_INDOOR}
    outdoor = {light.lower() for light in LIGHTS_OUTDOOR}
    for code in ("PAL-0001", "PAL-0002", "PAL-0007", "PAL-0042"):
        for variant in range(1, 9):
            scene = build_scene(_avatar_ctx(code, variant))
            body, _, light = scene.rpartition(". ")
            light = light.rstrip(".").lower()
            is_outdoor = any(
                text.lower() in body.lower() and flag for text, flag in CANDID_SCENES
            ) or any(
                word in body.lower()
                for word in ("rue", "marché", "parc", "arrêt de bus", "bord de l'eau",
                             "balcon", "voiture", "mur coloré", "escalier")
            )
            pool = outdoor if is_outdoor else indoor
            assert light in pool | indoor | outdoor
            if is_outdoor:
                assert light not in indoor - outdoor, f"{code}/{variant} : {scene}"


def test_scenes_spread_across_neighbouring_profiles():
    from backend.prompts.avatar_prompt import CANDID_SCENES, _varying_index

    picks = [_varying_index(len(CANDID_SCENES), _avatar_ctx(f"PAL-{n:04d}", 1), "candide")
             for n in range(1, 101)]
    assert not any(a == b for a, b in zip(picks, picks[1:]))


def test_hair_colour_agrees_grammatically():
    from backend.prompts.avatar_prompt import AvatarContext, build_avatar_prompt

    prompt = build_avatar_prompt(AvatarContext(
        code_femme="PAL-0007", age=28, ville="Dakar", pays="Sénégal",
        profession="Comptable", yeux="Noisette", cheveux="Châtain", variant_index=0,
    ))
    assert "de couleur" not in prompt
    assert "châtain" in prompt


def test_photo_style_switches_the_rendering_directive():
    from backend.prompts.avatar_prompt import AvatarContext, AvatarStyle, build_avatar_prompt

    common = dict(
        code_femme="PAL-0001", age=31, ville="Dakar", pays="Sénégal",
        profession="Comptable", yeux="Noir", cheveux="Noir", variant_index=1,
    )
    illustration = build_avatar_prompt(AvatarContext(**common, style=AvatarStyle.ILLUSTRATION))
    photo = build_avatar_prompt(AvatarContext(**common, style=AvatarStyle.PHOTO))
    assert "Illustration numérique" in illustration
    assert "prise au téléphone" in photo
    # Le rendu « photo » ne doit pas retomber dans l'esthétique studio, qui donne
    # le visage lissé que l'on repère immédiatement.
    for banni in ("studio", "50 mm", "retouche professionnelle"):
        assert banni not in photo
    assert "aucun lissage" in photo
    # Garde-fous communs aux deux rendus.
    for prompt in (illustration, photo):
        assert "seule sur l'image" in prompt
        assert "aucune personne réelle" in prompt


def test_download_path_is_constrained_to_valid_job_ids():
    """L'identifiant vient de l'URL et sert à bâtir un chemin : il doit être bridé."""
    from backend.api.routes import JOB_ID

    assert JOB_ID.match("JOB-20260814-70D1")
    for hostile in ("../../etc", "JOB-2026/../..", "JOB-20260814-70D1/../..", "", "JOB-x"):
        assert not JOB_ID.match(hostile), hostile


def _avatar(code: str, variant: int):
    from backend.prompts.avatar_prompt import AvatarContext

    return AvatarContext(
        code_femme=code, age=30, ville="Abidjan", pays="Côte d'Ivoire",
        profession="Infirmière", yeux="Marron", cheveux="Noir", variant_index=variant,
    )


def test_neighbouring_profiles_never_share_an_outfit():
    """Régression : un pas régulier entre codes retombait sur la même tenue.

    `seed` progressait de 8 exactement d'un profil au suivant et la liste des
    tenues comptait 8 entrées : toutes les femmes recevaient la même.
    """
    from backend.prompts.avatar_prompt import OUTFITS, build_outfit

    cuts = [build_outfit(_avatar(f"PAL-{n:04d}", 0)).split(" en ")[0].split(" dans ")[0]
            for n in range(1, 201)]
    assert not any(a == b for a, b in zip(cuts, cuts[1:])), "deux voisines habillées pareil"
    assert len(set(cuts[: len(OUTFITS)])) == len(OUTFITS), "le catalogue n'est pas parcouru"


def test_a_profile_never_repeats_an_outfit_across_its_photos():
    from backend.prompts.avatar_prompt import build_outfit

    for number in range(1, 51):
        tenues = {build_outfit(_avatar(f"PAL-{number:04d}", v)) for v in range(4)}
        assert len(tenues) == 4, f"PAL-{number:04d} répète une tenue"




def _fill(path: Path, codes: list[str], start_row: int, clear_example: bool) -> Path:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Femmes"]
    if clear_example:
        for column in range(1, 46):
            ws.cell(row=3, column=column).value = None
    for offset, code in enumerate(codes):
        ws.cell(row=start_row + offset, column=1, value=code)
    wb.save(path)
    return path


def test_blank_template_has_no_existing_profile(schema):
    """La ligne d'exemple porte PAL-0001 : la compter fausserait la reprise."""
    assert schema.femmes.example_row == 3
    assert schema.femmes.existing_codes == []
    assert schema.femmes.first_free_row == 3
    assert schema.femmes.next_number(CODE_FEMME) == 1


def test_numbering_resumes_after_the_last_existing_code(tmp_path):
    import shutil

    target = shutil.copy(TEMPLATE, tmp_path / "rempli.xlsx")
    _fill(Path(target), [f"PAL-{n:04d}" for n in range(1, 6)], 3, clear_example=True)

    filled = introspect(Path(target))
    assert filled.femmes.existing_codes == [f"PAL-{n:04d}" for n in range(1, 6)]
    assert filled.femmes.first_free_row == 8, "on écrirait par-dessus une ligne remplie"
    assert filled.femmes.next_number(CODE_FEMME) == 6


def test_data_kept_when_the_example_row_is_left_in_place(tmp_path):
    import shutil

    target = shutil.copy(TEMPLATE, tmp_path / "avec_exemple.xlsx")
    _fill(Path(target), ["PAL-0010", "PAL-0011"], 4, clear_example=False)

    filled = introspect(Path(target))
    assert filled.femmes.example_row == 3, "l'exemple doit rester reconnu"
    assert filled.femmes.existing_codes == ["PAL-0010", "PAL-0011"]
    assert filled.femmes.next_number(CODE_FEMME) == 12


def test_existing_codes_are_flagged_as_conflicts(schema):
    from backend.models.schemas import Profile
    from backend.services.validation_service import ValidationReport, Validator

    validator = Validator(schema, build_femme_specs(schema.femmes))
    report = ValidationReport()
    profile = Profile(
        code_femme="PAL-0003", code_agent="AG-01", nom_legal_complet="X",
        date_naissance="1995-01-01", nationalite="Ivoirienne",
        pays_residence="Côte d'Ivoire", ville_residence="Abidjan", prenom_affiche="X",
        ville_affichee="Abidjan", pays_affiche="Côte d'Ivoire", langues="Français",
        situation="Célibataire", enfants="Non", profession="Infirmière",
        niveau_etudes="Licence", taille_cm=165, poids_kg=60, yeux="Marron",
        cheveux="Noir", religion="Chrétienne", tabac="Non", alcool="Non",
        centres_interet="Cuisine", type_relation="Mariage", age_recherche_min=30,
        age_recherche_max=45, prete_a_demenager="Oui", accroche="a" * 50,
        presentation="b" * 700, recherche="c" * 250,
    )
    validator._check_against_existing([profile], ["PAL-0003"], report)
    assert report.errors, "une collision avec un code existant doit être une erreur"


class _FakeOpenAI:
    """Faux service : compte les appels et peut refuser la retouche."""

    def __init__(self, edit_fails: bool = False) -> None:
        self.edit_fails = edit_fails
        self.generated: list[str] = []
        self.edited: list[str] = []

    async def image_png(self, prompt: str) -> bytes:
        self.generated.append(prompt)
        return b"\x89PNG\r\n\x1a\n" + b"\x00" * 32

    async def image_png_from_reference(self, prompt: str, reference: bytes) -> bytes:
        if self.edit_fails:
            raise RuntimeError("retouche non servie par cette passerelle")
        assert reference, "la référence doit être transmise"
        self.edited.append(prompt)
        return b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


def _profile_for_images():
    from backend.models.schemas import Profile

    return Profile(
        code_femme="PAL-0001", code_agent="AG-01", nom_legal_complet="A B",
        date_naissance="1996-04-08", nationalite="Ivoirienne",
        pays_residence="Côte d'Ivoire", ville_residence="Abidjan", prenom_affiche="A",
        ville_affichee="Abidjan", pays_affiche="Côte d'Ivoire", langues="Français",
        situation="Célibataire", enfants="Non", profession="Infirmière",
        niveau_etudes="Licence", taille_cm=168, poids_kg=62, yeux="Marron",
        cheveux="Noir", religion="Chrétienne", tabac="Non", alcool="Non",
        centres_interet="Cuisine", type_relation="Mariage", age_recherche_min=30,
        age_recherche_max=45, prete_a_demenager="Oui", accroche="a" * 50,
        presentation="b" * 700, recherche="c" * 250,
    )


def _run_images(fake, per_profile: int, tmp_path):
    import asyncio

    from backend.config import Settings
    from backend.services.image_generator import ImageGenerator

    settings = Settings(OPENAI_API_KEY="k", OPENAI_IMAGE_STYLE="photo")
    generator = ImageGenerator.__new__(ImageGenerator)
    generator._openai = fake
    generator._semaphore = asyncio.Semaphore(4)
    generator._style = __import__(
        "backend.prompts.avatar_prompt", fromlist=["AvatarStyle"]
    ).AvatarStyle.PHOTO
    return asyncio.run(
        generator.generate_for_profiles(
            [_profile_for_images()], per_profile, tmp_path / "photos", date.today()
        )
    )


def test_photos_after_the_first_are_anchored_on_it(tmp_path):
    """Le visage dérivait sur huit images : les suivantes partent de la première."""
    fake = _FakeOpenAI()
    rows, failures = _run_images(fake, 8, tmp_path)

    assert len(rows) == 8 and not failures
    assert len(fake.generated) == 1, "seule la première photo est décrite de zéro"
    assert len(fake.edited) == 7, "les sept autres doivent partir de la référence"
    assert all("même visage" in prompt.lower() for prompt in fake.edited)


def test_images_still_produced_when_the_edit_endpoint_refuses(tmp_path):
    """Une passerelle sans retouche ne doit pas priver le dossier de photos."""
    fake = _FakeOpenAI(edit_fails=True)
    rows, failures = _run_images(fake, 5, tmp_path)

    assert len(rows) == 5, "le repli doit produire toutes les photos"
    assert not failures
    assert len(fake.generated) == 5, "chaque photo retombe sur une génération décrite"
